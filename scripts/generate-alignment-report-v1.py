#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成产品侧vs技术侧全量对标分析与方案衔接评估报告V1.0"""

from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
import os

doc = Document()

# 样式
style = doc.styles['Normal']
style.font.name = 'Microsoft YaHei'
style.font.size = Pt(10.5)
style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
style.paragraph_format.line_spacing = 1.5

# 辅助函数
def add_heading(text, level):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = 'Microsoft YaHei'
        run.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')

def add_para(text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = 'Microsoft YaHei'
    run.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    run.bold = bold
    return p

def add_table(headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    # 表头
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(9)
    # 数据
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.rows[ri + 1].cells[ci]
            cell.text = str(val)
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
    return table

# ====== 文档正文 ======

add_heading('产品侧vs技术侧全量对标分析与方案衔接评估报告', 0)

add_para('本文档基于产品侧最新设计方案（V1.0）与技术侧两份详细设计文档进行逐项对标，识别差异和衔接点，给出明确的调整建议和责任归属。')
add_para('')
add_para('技术侧文档来源：')
add_para('1. 《根因分析客群策略推荐一体化详细设计》（以下简称"一体化设计"）')
add_para('2. 《指标底座与全树根因分析能力改造详细设计》（以下简称"指标底座设计"）')

# ====== 第一章：根因库对标 ======

add_heading('一、根因库对标', 1)

add_heading('1.1 根因来源', 2)

add_table(
    ['对标项', '产品侧方案', '技术侧方案', '差异', '调整建议'],
    [
        ['根因来源类型', '专家经验、AI生成（2种）', 'MANUAL、STRATEGY、MODEL、AGENT（4种），来源优先级MANUAL>STRATEGY>MODEL>AGENT', '产品侧2种vs技术侧4种。技术侧STRATEGY=策略库配置、MODEL=模型归因、AGENT=AI智能体产出，粒度更细', '产品侧调整为4种来源：专家经验、策略库配置、模型归因、AI智能体产出。与置信度来源权重对齐，4种来源对应4个权重值'],
        ['来源权重', '专家经验=1.0、AI生成=0.75（2个权重）', 'STRATEGY=1.00、MANUAL=1.00、AGENT=0.75（3个权重，MODEL权重未定义）', '产品侧只有2个权重，技术侧有3个权重且MODEL权重缺失', '来源权重改为4个：专家经验(MANUAL)=1.0、策略库配置(STRATEGY)=1.0、模型归因(MODEL)=0.85（在1.0和0.75之间插值，需与技术侧确认）、AI智能体产出(AGENT)=0.75'],
    ]
)

add_heading('1.2 根因库运行数据指标', 2)

add_table(
    ['对标项', '产品侧方案', '技术侧方案', '差异', '调整建议'],
    [
        ['引用率', '引用率=调度单引用次数/推荐展示次数', '技术侧无直接对应指标。技术侧有hit_count（命中次数）和accuracy_rate（准确率）在mcd_strategy_cause_pattern表', '产品侧引用率是新增概念，技术侧hit_count/accuracy_rate是模式匹配维度', '保留产品侧引用率定义。hit_count和accuracy_rate在技术侧是模式匹配层面的统计，与产品侧"运营人员选用"维度不同。需技术侧在dispatch-srv新增根因被调度单引用的统计字段'],
        ['命中率', '命中率=变化区间匹配次数/推荐展示次数。匹配规则：异常方向一致且change_min<=实际变化幅度<=change_max', '技术侧mcd_strategy_cause_pattern表有change_min、change_max、anomaly_direction字段，以及hit_count、accuracy_rate字段。一体化设计11.2节根因融合规则中来源优先级用于合并，未定义命中率计算', '产品侧命中率公式与技术侧cause_pattern的change_min/change_max字段对齐，但技术侧未定义"变化区间匹配次数"的统计口径', '命中率定义保留。需技术侧在每次根因融合时记录：本次分析中指标实际变化幅度（feature_analysis_results_indicator.mom_rate/yoy_rate）是否落在cause_pattern的[change_min, change_max]区间内，并累计统计'],
        ['命中率实际变化幅度取值', '取自feature_analysis_results_indicator中该次分析记录的指标变化值，计算窗口与指标配置一致（月度指标取月环比、日频指标取日环比）', '指标底座设计3.7.2节：feature_analysis_results_indicator新增字段mom_rate（环比）、yoy_rate（同比）、change_direction（UP/DOWN/FLAT）', '完全对齐，产品侧取值来源与技术侧字段一一对应', '无需调整，数据来源确认可用'],
        ['置信度公式', '置信度=命中率x0.60+引用率x0.25+来源权重x0.15', '技术侧11.1节问题节点评分：problem_score=anomaly_scorex0.30+contribution_scorex0.35+business_weight_scorex0.20+confidence_scorex0.15。技术11.8节策略评分：match_score=root_cause_match_scorex0.35+segment_match_scorex0.30+source_weightx0.20+history_scorex0.15', '产品侧置信度是根因级评分，技术侧problem_score是指标节点级评分、match_score是策略匹配评分，三者不在同一层级', '产品侧置信度是根因自身的可信度评分，与技术侧problem_score（指标问题严重度）和match_score（策略推荐匹配度）是不同概念。保留产品侧定义。需技术侧在根因层面新增引用率和命中率统计，并按产品公式计算置信度'],
        ['置信度中的来源权重', '4种来源权重：专家经验=1.0、策略库配置=1.0、模型归因=0.85、AI智能体产出=0.75', '技术侧11.8节source_weight定义：STRATEGY=1.00、MANUAL=1.00、AGENT=0.75。MODEL权重未给出', '技术侧缺少MODEL权重值。技术侧优先级排序MANUAL>STRATEGY>MODEL>AGENT，MODEL排第三', '模型归因权重0.85为插值建议值（在1.0和0.75之间），需与技术侧确认后最终确定'],
    ]
)

add_heading('1.3 根因诊断条件', 2)

add_table(
    ['对标项', '产品侧方案', '技术侧方案', '差异', '调整建议'],
    [
        ['诊断条件结构', '指标编码+场景+异常方向+变化区间[change_min,change_max]+客户属性+市场事件', 'mcd_strategy_cause_pattern表：indicator_code、business_scene_code、anomaly_direction、change_min、change_max、dimension_tags_json、segment_codes_json', '产品侧"客户属性"和"市场事件"条件对应技术侧dimension_tags_json和segment_codes_json，但技术侧是JSON灵活字段', '保留产品侧前端结构化展示。技术侧JSON字段可承载产品侧的多种条件类型。一期命中率仅校验变化区间，客户属性和市场事件不参与自动匹配，与产品侧方案一致'],
        ['一期命中率范围', '仅校验异常方向+变化区间，客户属性和市场事件不参与自动匹配', '技术侧未定义命中率统计口径，但cause_pattern的dimension_tags_json和segment_codes_json一期暂不用于自动匹配', '对齐', '无需调整'],
    ]
)

add_heading('1.4 根因融合规则', 2)

add_table(
    ['对标项', '产品侧方案', '技术侧方案', '差异', '调整建议'],
    [
        ['根因融合', '产品侧根因库管理页面负责根因的新增/编辑/查看，融合逻辑由技术侧处理', '一体化设计11.2节：同一个问题指标下按描述相似度和根因ID合并，来源优先级MANUAL>STRATEGY>MODEL>AGENT，score取max(rootCauseScore)', '产品侧不涉及融合算法，但根因来源类型需要与融合优先级对齐', '产品侧根因来源改为4种（与融合优先级对齐），融合逻辑由技术侧实现，产品侧仅展示融合结果'],
    ]
)

# ====== 第二章：策略库对标 ======

add_heading('二、策略库对标', 1)

add_heading('2.1 策略来源与评分', 2)

add_table(
    ['对标项', '产品侧方案', '技术侧方案', '差异', '调整建议'],
    [
        ['策略来源类型', '专家经验、AI生成（2种）', '同根因来源：MANUAL、STRATEGY、MODEL、AGENT（4种）', '同根因库问题，产品侧2种vs技术侧4种', '策略来源也调整为4种：专家经验、策略库配置、模型归因、AI智能体产出'],
        ['策略匹配评分', '产品侧未定义策略推荐评分', '技术侧11.8节match_score=root_cause_match_scorex0.35+segment_match_scorex0.30+source_weightx0.20+history_scorex0.15', '技术侧有完整的策略推荐评分体系，产品侧未涉及', '策略推荐评分属于技术侧推荐引擎内部逻辑，产品侧不定义。但策略库运行数据的置信度需要与推荐评分中的source_weight对齐'],
        ['策略来源权重', '专家经验=1.0、AI生成=0.75', 'source_weight：STRATEGY=1.00、MANUAL=1.00、AGENT=0.75', '同根因库问题', '同根因库调整，策略来源权重改为4个'],
    ]
)

add_heading('2.2 策略库运行数据指标', 2)

add_table(
    ['对标项', '产品侧方案', '技术侧方案', '差异', '调整建议'],
    [
        ['转化率', '单调度单转化率=业务办理成功客户数/触达成功客户数；加权转化率=SUM(调度单转化率x触达成功客户数)/SUM(触达成功客户数)', '技术侧dispatch-srv负责工单调度和执行，9.4节有状态回传机制（dispatch_status=FINISHED），但未定义转化率统计口径', '技术侧有调度执行和状态回传，但未定义"业务办理成功"的统计闭环', '需技术侧补充：1.dispatch-srv记录触达成功客户数（渠道回传）；2.BSS/CRM回传业务办理结果；3.portal-srv聚合计算加权转化率'],
        ['转化率数据来源', '分子（业务办理成功客户数）：BSS/CRM业务办理结果回传；分母（触达成功客户数）：dispatch-srv触达结果统计', 'dispatch-srv 9.4节回传字段：dispatchStatus、customerGroupCode、customerCount，未含触达成功数和业务办理数', '技术侧当前回传字段不支撑转化率计算', '需技术侧dispatch-srv扩展回传字段：新增触达成功客户数(successReachCount)、业务办理成功客户数(bizSuccessCount)。或由portal-srv通过调度单关联BSS/CRM查询'],
        ['引用率', '引用率=调度单引用次数/推荐展示次数', '技术侧无直接对应。策略推荐展示在portal-web 10.5节策略推荐面板，引用（采纳）操作也在此', '产品侧引用率需要在技术侧策略推荐链路中新增统计点', '需技术侧在策略推荐面板新增统计：1.策略被推荐展示次数（已有推荐行为）；2.运营人员采纳策略创建调度单次数（需在采纳操作时记录）'],
        ['置信度公式', '置信度=转化率x0.70+引用率x0.15+来源权重x0.15', '技术侧无根因/策略级置信度定义。技术侧11.1节confidence_score是指标数据可信分，与产品侧置信度不同', '产品侧策略置信度是策略自身的可信度，技术侧confidence_score是指标数据质量评分，两者概念不同', '保留产品侧定义。需技术侧在策略层面新增引用率和转化率统计，并按产品公式计算策略置信度'],
    ]
)

add_heading('2.3 策略动作类型', 2)

add_table(
    ['对标项', '产品侧方案', '技术侧方案', '差异', '调整建议'],
    [
        ['策略动作类型', '产品侧策略编辑页7层Tab中第5层"推荐产品/方案"含策略动作类型', '技术侧8.4.3节：mcd_strategy_info新增action_type字段', '产品侧和技术侧均有策略动作类型，字段对齐', '无需调整'],
    ]
)

# ====== 第三章：指标底座对标 ======

add_heading('三、指标底座对标', 1)

add_heading('3.1 指标分析配置', 2)

add_table(
    ['对标项', '产品侧方案', '技术侧方案（指标底座设计）', '差异', '调整建议'],
    [
        ['指标参与分析', '指标库配置页：是否参与分析', 'indicator_info新增participate_analysis字段', '完全对齐', '无需调整'],
        ['指标参与TopN', '指标库配置页：是否参与TopN排序', 'indicator_info新增participate_topn字段', '完全对齐', '无需调整'],
        ['指标业务权重', '指标库配置页：默认业务权重', 'indicator_info新增default_business_weight字段', '完全对齐', '无需调整'],
        ['指标分析模式', '指标库配置页：默认分析模式（规则/相关性/模型/混合）', 'indicator_info新增default_analysis_mode字段，取值RULE_FORMULA/CORRELATION/SHAP_MODEL/HYBRID', '完全对齐', '无需调整'],
        ['指标方向', '指标库配置页：指标方向（正向/反向）', '复用现有indicator_direction字段，节点级差异通过indicator_node_analysis_override覆盖', '完全对齐', '无需调整'],
        ['指标关系类型', '指标关系配置：关系类型（树/求和/比值/乘积/引用）', 'indicator_relation新增relation_type字段', '完全对齐', '无需调整'],
        ['指标关系方向', '指标关系配置：关系方向（正相关/负相关/未知）', 'indicator_relation新增relation_direction字段', '完全对齐', '无需调整'],
    ]
)

add_heading('3.2 指标树/指标组分析配置', 2)

add_table(
    ['对标项', '产品侧方案', '技术侧方案（指标底座设计）', '差异', '调整建议'],
    [
        ['指标树分析开关', '指标树配置页：是否启用根因分析', 'indicator_tree_info新增analysis_enabled字段', '完全对齐', '无需调整'],
        ['指标树TopN数量', '指标树配置页：TopN数量（默认5）', 'indicator_tree_info新增topn_count字段，默认5', '完全对齐', '无需调整'],
        ['指标树评分权重', '指标树配置页：综合评分权重配置', 'indicator_tree_info新增problem_score_config_json字段', '完全对齐', '无需调整'],
        ['指标组分析模式', '指标组配置页：分析模式（规则/相关性/模型/混合）', 'indicator_group_info新增analysis_mode字段', '完全对齐', '无需调整'],
        ['指标组训练触发', '指标组配置页：训练触发方式（手动/发布时自动）', 'indicator_group_info新增training_trigger_mode字段', '完全对齐', '无需调整'],
        ['指标组降级策略', '指标组配置页：模型不可用时降级策略', 'indicator_group_info新增degrade_strategy字段', '完全对齐', '无需调整'],
        ['节点级覆盖配置', '指标树节点：个别节点差异配置（不参与分析/方向覆盖等）', '新增indicator_node_analysis_override表', '完全对齐', '无需调整'],
    ]
)

add_heading('3.3 分析结果与评分', 2)

add_table(
    ['对标项', '产品侧方案', '技术侧方案（指标底座设计）', '差异', '调整建议'],
    [
        ['问题指标评分', '指标树分析结果：综合评分排序+Top5标识', 'problem_score=anomaly_scorex0.30+contribution_scorex0.35+business_weight_scorex0.20+confidence_scorex0.15', '完全对齐', '无需调整。注意：此处confidence_score是指标数据可信分，不是根因/策略置信度'],
        ['根因融合展示', '指标分析面板：展示agent根因和strategy根因的合并结果', '一体化设计11.2节：按描述相似度和根因ID合并，来源优先级MANUAL>STRATEGY>MODEL>AGENT', '对齐，但产品侧需展示4种来源', '产品侧根因列表展示来源时需与技术侧4种来源对齐'],
        ['分析结果字段', '分析面板展示异动强度、贡献度、业务权重、数据可信度', 'feature_analysis_results_indicator新增anomaly_score、contribution_score、business_weight_score、confidence_score等字段', '完全对齐', '无需调整'],
    ]
)

# ====== 第四章：微服务衔接评估 ======

add_heading('四、微服务衔接评估', 1)

add_table(
    ['微服务', '产品侧需求', '技术侧现有能力', '缺口', '衔接建议'],
    [
        ['portal-srv', '1.根因/策略的引用率、命中率、转化率统计\n2.置信度计算\n3.根因/策略列表、详情、新增/编辑页面', '1.分析结果聚合\n2.指标配置管理\n3.根因融合\n4.Top5评分', '1.缺少引用率统计（需记录推荐展示次数和调度单引用次数）\n2.缺少命中率统计（需在根因融合时记录变化区间匹配结果）\n3.缺少转化率统计（需聚合dispatch-srv触达结果和BSS/CRM办理结果）', '1.portal-srv新增根因/策略推荐日志表，记录每次推荐展示\n2.根因融合逻辑中新增命中率统计\n3.新增转化率聚合计算服务'],
        ['strategy-srv', '1.根因/策略的CRUD\n2.根因来源字段（4种来源）\n3.根因模式匹配（change_min/change_max）', '1.已有batchQueryRootCause、batchQueryStrategy接口\n2.mcd_strategy_cause_pattern表含change_min/change_max/anomaly_direction\n3.cause_info表已有confidence_score字段', '1.来源字段需从2种扩展到4种（MANUAL/STRATEGY/MODEL/AGENT）\n2.hit_count和accuracy_rate字段已存在但口径需与产品侧命中率对齐', '1.扩展来源字段为4种枚举\n2.hit_count复用为"变化区间匹配次数"，accuracy_rate复用为"命中率"'],
        ['dispatch-srv', '1.触达成功客户数统计\n2.业务办理成功客户数统计（或BSS/CRM回传）\n3.调度单引用根因/策略关联', '1.已有客群调度工单\n2.9.4节有状态回传（dispatchStatus、customerCount）\n3.工单已有rootCauseId关联', '1.回传字段缺少触达成功数和业务办理成功数\n2.需新增策略关联字段', '1.dispatch-srv回传扩展successReachCount和bizSuccessCount\n2.工单新增strategyId关联字段'],
        ['ai-agent-srv', '1.根因分析时记录推荐展示次数\n2.策略推荐时记录推荐展示次数', '1.已承担SHAP/LLM/客群/策略分析编排\n2.已有回写portal分析结果能力', '1.推荐展示行为未记录', 'agent回写分析结果时新增推荐展示计数'],
        ['ai-train-srv', '无直接需求', '提供SHAP分析、相关性分析', '无缺口', '无需调整'],
    ]
)

# ====== 第五章：数据衔接评估 ======

add_heading('五、数据衔接评估', 1)

add_heading('5.1 根因库运行数据衔接', 2)

add_table(
    ['指标', '分子取值', '分母取值', '数据来源服务', '当前技术侧是否支撑', '衔接方案'],
    [
        ['引用率', '调度单引用次数：从dispatch-srv调度单的rootCauseId关联字段统计', '推荐展示次数：portal-srv根因融合时该根因被召回展示的次数', 'dispatch-srv+portal-srv', '部分支撑。调度单已有rootCauseId关联，但推荐展示次数未记录', 'portal-srv新增根因推荐日志表root_cause_recommend_log，每次根因融合时插入一条记录'],
        ['命中率', '变化区间匹配次数：根因融合时指标实际变化幅度落在[change_min,change_max]区间内的次数', '推荐展示次数（同引用率分母）', 'portal-srv（根因融合逻辑+feature_analysis_results_indicator）', '部分支撑。cause_pattern有change_min/change_max字段，feature_analysis_results_indicator有mom_rate/yoy_rate，但匹配统计未实现', '在根因融合逻辑RootCauseMergeService中新增命中率统计：每次融合时比较实际变化幅度与[change_min,change_max]，记录匹配结果'],
        ['置信度', '由命中率和引用率按公式计算', '由命中率和引用率按公式计算', 'portal-srv计算', '不支撑。需要命中率和引用率先实现', '在portal-srv新增RootCauseConfidenceService，按公式计算置信度并更新strategy-srv的cause_info.confidence_score'],
    ]
)

add_heading('5.2 策略库运行数据衔接', 2)

add_table(
    ['指标', '分子取值', '分母取值', '数据来源服务', '当前技术侧是否支撑', '衔接方案'],
    [
        ['转化率（单调度单）', '业务办理成功客户数：BSS/CRM在客户触达后7天内回传的办理成功记录', '触达成功客户数：dispatch-srv渠道执行后回传的触达成功记录', 'dispatch-srv+BSS/CRM', '不支撑。dispatch-srv回传字段不包含触达成功数和办理成功数', '1.dispatch-srv扩展回传字段：successReachCount、bizSuccessCount\n2.或portal-srv通过调度单关联BSS/CRM查询办理结果'],
        ['转化率（加权）', 'SUM(调度单转化率x触达成功客户数)', 'SUM(触达成功客户数)', 'portal-srv聚合计算', '不支撑。需先实现单调度单转化率', 'portal-srv新增StrategyConversionService，聚合计算加权转化率'],
        ['引用率', '调度单引用次数：从dispatch-srv调度单的strategyId关联字段统计', '推荐展示次数：portal-srv策略推荐时该策略被推荐展示的次数', 'dispatch-srv+portal-srv', '部分支撑。调度单需新增strategyId关联字段，推荐展示次数未记录', '1.dispatch-srv工单新增strategyId字段\n2.portal-srv新增策略推荐日志表strategy_recommend_log'],
        ['置信度', '由转化率和引用率按公式计算', '由转化率和引用率按公式计算', 'portal-srv计算', '不支撑。需要转化率和引用率先实现', '在portal-srv新增StrategyConfidenceService，按公式计算置信度'],
    ]
)

# ====== 第六章：关键差异与决策项 ======

add_heading('六、关键差异与决策项', 1)

add_table(
    ['序号', '差异项', '产品侧', '技术侧', '影响', '建议', '决策方'],
    [
        ['1', '根因/策略来源类型数量', '2种（专家经验、AI生成）', '4种（MANUAL、STRATEGY、MODEL、AGENT）', '来源类型影响置信度计算中的来源权重，进而影响根因融合优先级', '产品侧调整为4种来源，与技术侧对齐。前端新增/编辑页的"来源"下拉从2个选项扩展为4个', '产品侧确认'],
        ['2', '模型归因(MODEL)来源权重', '未定义（产品侧原来归入AI生成=0.75）', '未定义（技术侧source_weight只有STRATEGY/MANUAL/AGENT三个值）', '如果模型归因与AI智能体同权重0.75，则无需区分；如果不同权重，则需明确', '建议模型归因权重=0.85（介于1.0和0.75之间），需技术侧确认', '技术侧确认'],
        ['3', '转化率数据闭环', '产品侧定义了转化率公式，依赖BSS/CRM回传', '技术侧dispatch-srv当前回传字段不含触达成功数和办理成功数', '转化率无法计算，策略库置信度为空', '一期可先展示"暂无数据"，二期实现BSS/CRM办理结果回传闭环', '双方协商'],
        ['4', '命中率统计口径', '产品侧定义了命中率=变化区间匹配次数/推荐展示次数', '技术侧cause_pattern有hit_count和accuracy_rate字段，但统计口径未与产品侧对齐', '命中率字段已有但口径可能不一致', '复用hit_count字段，按产品侧口径重新定义统计逻辑：在根因融合时比对实际变化幅度与[change_min,change_max]', '技术侧确认'],
        ['5', '推荐展示次数统计', '产品侧引用率的分母需要"推荐展示次数"', '技术侧未记录每次推荐/召回行为', '引用率无法计算', 'portal-srv新增推荐日志表，记录每次根因/策略被推荐展示的行为', '技术侧实现'],
        ['6', '策略关联调度单', '产品侧引用率分子需要"调度单引用次数"', '技术侧工单已有rootCauseId关联，但缺少strategyId关联', '策略引用率分子无法统计', 'dispatch-srv工单新增strategyId关联字段', '技术侧实现'],
    ]
)

# ====== 第七章：分阶段衔接建议 ======

add_heading('七、分阶段衔接建议', 1)

add_heading('7.1 一阶段（最小可用）', 2)

add_table(
    ['能力', '产品侧', '技术侧', '状态'],
    [
        ['根因/策略来源字段扩展为4种', '前端下拉选项扩展', 'strategy-srv扩展source字段枚举', '需实施'],
        ['根因命中率统计', '展示命中率卡片', 'portal-srv在根因融合时统计变化区间匹配次数和推荐展示次数', '需实施'],
        ['根因引用率统计', '展示引用率卡片', 'portal-srv记录根因推荐展示次数+dispatch-srv统计调度单引用次数', '需实施'],
        ['根因置信度计算', '展示置信度卡片', 'portal-srv按公式计算', '需实施'],
        ['策略引用率统计', '展示引用率卡片', 'dispatch-srv工单新增strategyId+portal-srv记录策略推荐展示次数', '需实施'],
        ['策略转化率', '展示"暂无数据"', '二期实现BSS/CRM回传闭环', '一期不实施'],
        ['策略置信度', '转化率暂无数据时置信度也暂无', '依赖转化率', '一期不实施'],
    ]
)

add_heading('7.2 二阶段（数据闭环）', 2)

add_table(
    ['能力', '说明', '依赖'],
    [
        ['策略转化率统计', 'dispatch-srv回传触达成功数和BSS/CRM回传业务办理成功数', 'BSS/CRM接口对接'],
        ['策略置信度计算', '基于转化率和引用率按公式计算', '转化率统计实现'],
        ['模型归因来源权重确认', '与技术侧确认MODEL权重值（建议0.85）', '技术侧评估'],
        ['命中率统计增强', '除变化区间外，增加客户属性和市场事件条件的自动匹配', 'CDP实时查询能力+市场事件维护机制'],
    ]
)

# 保存
output_path = '/workspace/projects/docs/对标分析/产品侧vs技术侧全量对标分析与方案衔接评估报告.docx'
os.makedirs(os.path.dirname(output_path), exist_ok=True)
doc.save(output_path)
print(f'报告已保存到: {output_path}')

# 验证
doc2 = Document(output_path)
table_count = len(doc2.tables)
cell_count = 0
empty_count = 0
for t in doc2.tables:
    for row in t.rows:
        for cell in row.cells:
            cell_count += 1
            if not cell.text.strip():
                empty_count += 1
print(f'表格数: {table_count}, 单元格数: {cell_count}, 空单元格数: {empty_count}')
