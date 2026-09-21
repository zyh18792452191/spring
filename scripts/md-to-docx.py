#!/usr/bin/env python3
"""将 PRD-流量运营区域攻坚.md 转换为 Word 文档"""
import os, re
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MD_PATH = os.path.join(BASE_DIR, 'docs', 'PRD文档', 'PRD-流量运营区域攻坚.md')
IMG_DIR = os.path.join(BASE_DIR, 'docs', 'PRD文档', 'screenshots')
OUTPUT_PATH = os.path.join(BASE_DIR, 'docs', 'PRD文档', 'PRD-流量运营区域攻坚.docx')

# ─── Helper Functions ────────────────────────────────────────────────

def set_cell_shading(cell, color_hex):
    shading = OxmlElement('w:shd')
    shading.set(qn('w:val'), 'clear')
    shading.set(qn('w:color'), 'auto')
    shading.set(qn('w:fill'), color_hex)
    cell._tc.get_or_add_tcPr().append(shading)

def add_run_with_format(paragraph, text, bold=False, size=Pt(10.5), color=None):
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = size
    if color:
        run.font.color.rgb = color
    return run

def parse_inline(para_element, text, base_size=Pt(10.5)):
    """Parse inline markdown formatting (bold, code, links, images) into docx runs."""
    # Pattern: **bold**, `code`, [text](url), ![alt](path)
    pattern = re.compile(r'(\*\*(.+?)\*\*)|(`(.+?)`)|(\[(.+?)\]\((.+?)\))|(!\[(.+?)\]\((.+?)\))')
    last_end = 0
    for m in pattern.finditer(text):
        # Add plain text before this match
        if m.start() > last_end:
            plain = text[last_end:m.start()]
            add_run_with_format(para_element, plain, size=base_size)
        
        if m.group(1):  # **bold**
            add_run_with_format(para_element, m.group(2), bold=True, size=base_size, color=RGBColor(0, 0, 0))
        elif m.group(3):  # `code`
            run = add_run_with_format(para_element, m.group(4), size=base_size, color=RGBColor(0, 0, 0))
            run.font.name = 'Consolas'
            run.font.size = Pt(9)
        elif m.group(5):  # [text](url) - link
            add_run_with_format(para_element, m.group(6), size=base_size, color=RGBColor(0, 0, 0))
        elif m.group(8):  # ![alt](path) - image
            img_path = m.group(10)
            # Resolve relative path
            if not os.path.isabs(img_path):
                full_img = os.path.join(BASE_DIR, 'docs', 'PRD文档', img_path)
            else:
                full_img = img_path
            if os.path.exists(full_img):
                para_element.add_run('\n')
                run = para_element.add_run()
                run.add_picture(full_img, width=Inches(6.0))
            else:
                add_run_with_format(para_element, f'[图片: {m.group(9)}]', size=base_size, color=RGBColor(0, 0, 0))
        
        last_end = m.end()
    
    # Remaining plain text
    if last_end < len(text):
        add_run_with_format(para_element, text[last_end:], size=base_size)

def add_image_to_doc(doc, img_path, width_inches=6.0, caption=None):
    """Add an image with optional caption."""
    if os.path.exists(img_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(img_path, width=Inches(width_inches))
        if caption:
            cap = doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = cap.add_run(caption)
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(0, 0, 0)
    else:
        p = doc.add_paragraph(f'[图片未找到: {img_path}]')

# ─── Main Conversion ────────────────────────────────────────────────

def convert():
    with open(MD_PATH, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = Document()
    
    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Microsoft YaHei'
    font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')

    i = 0
    first_heading = True
    while i < len(lines):
        line = lines[i].rstrip('\n')

        # Skip empty lines
        if not line.strip():
            i += 1
            continue

        # Horizontal rule
        if line.strip() == '---':
            # Add a thin line
            p = doc.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            pBdr = OxmlElement('w:pBdr')
            bottom = OxmlElement('w:bottom')
            bottom.set(qn('w:val'), 'single')
            bottom.set(qn('w:sz'), '4')
            bottom.set(qn('w:space'), '1')
            bottom.set(qn('w:color'), 'CCCCCC')
            pBdr.append(bottom)
            pPr.append(pBdr)
            i += 1
            continue

        # Headings
        heading_match = re.match(r'^(#{1,6})\s+(.+)$', line)
        if heading_match:
            level = len(heading_match.group(1))
            text = heading_match.group(2).strip()
            # Map markdown heading levels to Word heading levels
            # ## = Heading 1 (top level in doc), ### = Heading 2, etc.
            word_level = max(1, level - 1)
            h = doc.add_heading(text, level=word_level)
            # Center the first heading (document title)
            if first_heading:
                h.alignment = WD_ALIGN_PARAGRAPH.CENTER
                first_heading = False
            i += 1
            continue

        # Image line: ![alt](path)
        img_match = re.match(r'^!\[(.+?)\]\((.+?)\)\s*$', line)
        if img_match:
            alt = img_match.group(1)
            img_rel = img_match.group(2)
            if not os.path.isabs(img_rel):
                full_img = os.path.join(BASE_DIR, 'docs', 'PRD文档', img_rel)
            else:
                full_img = img_rel
            add_image_to_doc(doc, full_img, caption=alt)
            i += 1
            continue

        # Table
        if '|' in line and i + 1 < len(lines) and re.match(r'^\|[\s\-:|]+\|$', lines[i+1].strip()):
            # Parse table header
            headers = [c.strip() for c in line.strip().strip('|').split('|')]
            i += 1  # skip separator line
            i += 1  # move to first data row
            
            rows = []
            while i < len(lines) and '|' in lines[i] and lines[i].strip().startswith('|'):
                row_data = [c.strip() for c in lines[i].strip().strip('|').split('|')]
                rows.append(row_data)
                i += 1
            
            # Create table
            num_cols = len(headers)
            table = doc.add_table(rows=1 + len(rows), cols=num_cols)
            table.style = 'Table Grid'
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            
            # Header row
            for ci, h in enumerate(headers):
                cell = table.rows[0].cells[ci]
                cell.text = ''
                p = cell.paragraphs[0]
                r = p.add_run(h)
                r.bold = True
                r.font.size = Pt(10)
                set_cell_shading(cell, 'D9E2F3')
            
            # Data rows
            for ri, row in enumerate(rows):
                for ci in range(min(num_cols, len(row))):
                    cell = table.rows[ri + 1].cells[ci]
                    cell.text = ''
                    p = cell.paragraphs[0]
                    val = row[ci]
                    # Handle inline formatting in table cells
                    parse_inline(p, val, base_size=Pt(10))
            
            # Add small spacing after table
            doc.add_paragraph()
            continue

        # Numbered list (e.g., "1. xxx")
        num_list_match = re.match(r'^(\d+)\.\s+(.+)$', line)
        if num_list_match:
            p = doc.add_paragraph(style='List Number')
            parse_inline(p, num_list_match.group(2))
            i += 1
            continue

        # Bullet list (e.g., "- xxx" or "* xxx")
        bullet_match = re.match(r'^[\-\*]\s+(.+)$', line)
        if bullet_match:
            p = doc.add_paragraph(style='List Bullet')
            parse_inline(p, bullet_match.group(1))
            i += 1
            continue

        # Indented sub-bullet (e.g., "  - xxx")
        sub_bullet_match = re.match(r'^\s{2,}[\-\*]\s+(.+)$', line)
        if sub_bullet_match:
            p = doc.add_paragraph(style='List Bullet 2')
            parse_inline(p, sub_bullet_match.group(1))
            i += 1
            continue

        # Code block
        if line.strip().startswith('```'):
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith('```'):
                code_lines.append(lines[i].rstrip('\n'))
                i += 1
            i += 1  # skip closing ```
            
            code_text = '\n'.join(code_lines)
            p = doc.add_paragraph()
            p.style = doc.styles['Normal']
            # Add left border for code block
            pPr = p._p.get_or_add_pPr()
            pBdr = OxmlElement('w:pBdr')
            left_bdr = OxmlElement('w:left')
            left_bdr.set(qn('w:val'), 'single')
            left_bdr.set(qn('w:sz'), '12')
            left_bdr.set(qn('w:space'), '4')
            left_bdr.set(qn('w:color'), 'CCCCCC')
            pBdr.append(left_bdr)
            pPr.append(pBdr)
            
            run = p.add_run(code_text)
            run.font.name = 'Consolas'
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0, 0, 0)
            continue

        # Regular paragraph
        p = doc.add_paragraph()
        parse_inline(p, line.strip())
        i += 1

    # Save
    os.makedirs(os.path.dirname(OUTPUT_PATH), exist_ok=True)
    doc.save(OUTPUT_PATH)
    print(f'Word文档已生成: {OUTPUT_PATH}')

if __name__ == '__main__':
    convert()
