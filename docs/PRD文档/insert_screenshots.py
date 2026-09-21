#!/usr/bin/env python3
"""在已生成的PRD Word文档中插入截图"""

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import os

SCREENSHOT_DIR = '/workspace/projects/docs/PRD文档/screenshots'
DOC_PATH = '/workspace/projects/docs/PRD文档/AI洞察报告PRD V3.0.docx'
OUTPUT_PATH = '/workspace/projects/docs/PRD文档/AI洞察报告PRD V3.0.docx'
FONT_NAME = '微软雅黑'

def add_image_after_element(doc, heading_elem, filepath, caption, width_inches=5.2):
    """在标题元素后插入截图+标题"""
    if not os.path.exists(filepath):
        # 添加占位文字
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f'[截图待补充: {caption}]')
        run.font.size = Pt(10)
        run.font.italic = True
        run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
        p_elem = p._element
        p_elem.getparent().remove(p_elem)
        heading_elem.addnext(p_elem)
        return
    
    # 截图段落
    pic_para = doc.add_paragraph()
    pic_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = pic_para.add_run()
    run.add_picture(filepath, width=Inches(width_inches))
    pic_elem = pic_para._element
    pic_elem.getparent().remove(pic_elem)
    heading_elem.addnext(pic_elem)
    
    # 标题段落
    cap_para = doc.add_paragraph()
    cap_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap_run = cap_para.add_run(f'图：{caption}')
    cap_run.font.size = Pt(9)
    cap_run.font.italic = True
    cap_run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    cap_run.font.name = FONT_NAME
    cap_run._element.rPr.rFonts.set(qn('w:eastAsia'), FONT_NAME)
    cap_elem = cap_para._element
    cap_elem.getparent().remove(cap_elem)
    # 插到截图前面（即标题后面）
    heading_elem.addnext(cap_elem)

def insert_screenshots():
    doc = Document(DOC_PATH)
    
    # 收集所有标题
    heading_elems = {}
    for para in doc.paragraphs:
        if para.style.name.startswith('Heading'):
            heading_elems[para.text] = para._element
    
    # 截图映射 - key匹配文档中的实际标题文本
    screenshot_plan = [
        # (标题文本片段, 截图文件名, 截图标题)
        ('3.1 AI洞察助手', 'ai-insight-main.png', 'AI洞察助手主页面'),
        ('3.2 数据源管理', 'datasource-list.png', '数据源管理页面'),
        ('3.3 模板管理', 'template-list.png', '模板管理页面'),
        ('3.4 下载专区', 'download-zone-ai-tab.png', '下载专区-AI洞察分析Tab'),
    ]
    
    # 额外的弹窗截图：在3.1标题后追加（多张）
    ai_extra_screenshots = [
        ('ai-insight-start-page.png', 'AI洞察助手开始页-快捷入口卡片'),
        ('ai-insight-data-preparation.png', '数据准备阶段-模板选择'),
        ('ai-insight-datasource-buttons.png', '跳过模板后-数据源按钮区'),
        ('ai-insight-indicator-dialog.png', '平台指标选择弹窗'),
        ('ai-insight-dataset-dialog.png', '数据集选择弹窗'),
        ('ai-insight-group-dialog.png', '客群选择弹窗'),
        ('ai-insight-report-dialog.png', '报表选择弹窗'),
        ('ai-insight-tree-dialog.png', '指标树选择弹窗'),
        ('ai-insight-selected-datasources.png', '已选数据源列表+账期选择器'),
        ('ai-insight-followup-stage.png', '追问深化阶段'),
    ]
    
    # 数据源管理额外截图
    ds_extra = [
        ('datasource-detail-dialog.png', '数据源详情弹窗'),
    ]
    
    # 模板管理额外截图
    tpl_extra = [
        ('template-ai-create.png', 'AI对话创建模板页面'),
        ('template-manual-edit.png', '手动创建模板页面'),
        ('template-extract.png', '从历史报告提取模板页面'),
    ]
    
    # 下载专区额外截图
    dl_extra = [
        ('download-zone.png', '下载专区页面-全部Tab'),
    ]
    
    # 执行插入（从后往前，避免索引偏移）
    # 先按标题分组
    insert_groups = [
        ('3.4 下载专区', dl_extra + [('download-zone-ai-tab.png', '下载专区-AI洞察分析Tab')]),
        ('3.3 模板管理', tpl_extra),
        ('3.2 数据源管理', ds_extra),
        ('3.1 AI洞察助手', ai_extra_screenshots),
    ]
    
    total_inserted = 0
    for heading_text_frag, screenshots in insert_groups:
        # 找到匹配的标题
        heading_elem = None
        for text, elem in heading_elems.items():
            if heading_text_frag in text:
                heading_elem = elem
                break
        
        if not heading_elem:
            print(f'[WARN] 未找到标题: {heading_text_frag}')
            continue
        
        # 反向插入（从后往前），这样每张截图都紧接标题后面
        for filepath, caption in reversed(screenshots):
            full_path = os.path.join(SCREENSHOT_DIR, filepath)
            add_image_after_element(doc, heading_elem, full_path, caption)
            total_inserted += 1
            print(f'  插入: {caption}')
    
    doc.save(OUTPUT_PATH)
    print(f'\n共插入 {total_inserted} 张截图')
    print(f'文档已保存: {OUTPUT_PATH}')
    fsize = os.path.getsize(OUTPUT_PATH)
    print(f'文件大小: {fsize/1024:.1f} KB')

if __name__ == '__main__':
    insert_screenshots()
